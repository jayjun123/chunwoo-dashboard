import openpyxl
import pandas as pd

try:
    # 엑셀 파일 로드
    wb = openpyxl.load_workbook('public/통합 문서1 8월.xlsx')
    print('시트 목록:', wb.sheetnames)
    
    # 첫 번째 시트 선택
    ws = wb[wb.sheetnames[0]]
    print(f'첫 번째 시트: {wb.sheetnames[0]}')
    
    # A1:D20 데이터 확인
    print('\nA1:D20 데이터:')
    for row in range(1, 21):
        row_data = []
        for col in range(1, 5):
            cell_value = ws.cell(row=row, column=col).value
            row_data.append(str(cell_value) if cell_value is not None else '')
        print(f'행 {row:2d}: {row_data}')
    
    # 현장명이 있을 것 같은 행들 더 자세히 확인
    print('\n현장명이 있을 것 같은 행들 (C열 중심):')
    for row in range(1, 51):
        c_value = ws.cell(row=row, column=3).value
        if c_value and str(c_value).strip():
            print(f'행 {row:2d} C열: "{c_value}"')
    
except Exception as e:
    print(f'오류 발생: {e}')
